from pathlib import Path


def export_minutes(minutes, output_path, file_format=None):
    """
    Export structured meeting minutes to TXT or DOCX.
    minutes should contain topics, decisions, and action_items.
    """
    output_path = Path(output_path)
    fmt = (file_format or output_path.suffix.lstrip(".")).lower()

    if fmt not in ("txt", "docx"):
        raise ValueError("Supported formats are TXT and DOCX.")

    lines = ["MEETING MINUTES", ""]

    lines.append("DISCUSSION TOPICS")
    lines.extend(
        f"- {item}" for item in minutes.get("topics", [])
    )
    lines.append("")

    lines.append("DECISIONS")
    lines.extend(
        f"- {item}" for item in minutes.get("decisions", [])
    )
    lines.append("")

    lines.append("ACTION ITEMS")
    for item in minutes.get("action_items", []):
        task = item.get("task") or "Not specified"
        owner = item.get("owner") or "Not assigned"
        deadline = item.get("deadline") or "Not specified"

        lines.append(f"- Task: {task}")
        lines.append(f"  Owner: {owner}")
        lines.append(f"  Deadline: {deadline}")

    if fmt == "txt":
        output_path.write_text(
            "\n".join(lines) + "\n",
            encoding="utf-8",
        )
    else:
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError(
                "Install python-docx using: "
                "python3 -m pip install python-docx"
            ) from exc

        document = Document()
        document.add_heading("Meeting Minutes", level=0)

        for section, items in (
            ("Discussion Topics", minutes.get("topics", [])),
            ("Decisions", minutes.get("decisions", [])),
        ):
            document.add_heading(section, level=1)
            if items:
                for item in items:
                    document.add_paragraph(str(item), style="List Bullet")
            else:
                document.add_paragraph("None recorded.")

        document.add_heading("Action Items", level=1)
        action_items = minutes.get("action_items", [])
        if action_items:
            for item in action_items:
                document.add_paragraph(
                    f"Task: {item.get('task') or 'Not specified'}"
                )
                document.add_paragraph(
                    f"Owner: {item.get('owner') or 'Not assigned'}"
                )
                document.add_paragraph(
                    f"Deadline: {item.get('deadline') or 'Not specified'}"
                )
        else:
            document.add_paragraph("None recorded.")

        document.save(str(output_path))

    return str(output_path)
